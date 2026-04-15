import customtkinter as ctk
import random
import math
import tkinter.filedialog as fd
import openpyxl
import os
import sys
import tempfile
import shutil

# Настройка внешнего вида
ctk.set_appearance_mode("Dark")
ctk.set_default_color_theme("blue")

def resource_path(relative_path):
    """ Получает абсолютный путь к ресурсу, работает для dev и для PyInstaller """
    try:
        # PyInstaller создает временную папку и сохраняет путь в _MEIPASS
        base_path = sys._MEIPASS
    except Exception:
        base_path = os.path.abspath(".")

    return os.path.join(base_path, relative_path)

class EntropyCalculatorApp(ctk.CTk):
    def __init__(self):
        super().__init__()

        self.title("Entro.py")
        self.geometry("800x600") 
        
        # === НАСТРОЙКА ИКОНКИ ВНУТРИ ОКНА ===
        try:
            # 1. Пытаемся найти иконку в ресурсах PyInstaller
            if getattr(sys, 'frozen', False):
                # Если запущено как .exe
                icon_src = resource_path("entro.ico")
            else:
                # Если запущено как скрипт
                icon_src = os.path.join(os.path.dirname(__file__), "entro.ico")

            # 2. Для надежности копируем иконку во временную папку, 
            # так как некоторые версии Windows/Tkinter плохо читают файлы из temp архива
            with tempfile.NamedTemporaryFile(delete=False, suffix='.ico') as tmp_file:
                with open(icon_src, 'rb') as src_file:
                    shutil.copyfileobj(src_file, tmp_file)
                temp_icon_path = tmp_file.name

            # 3. Устанавливаем иконку
            self.iconbitmap(temp_icon_path)
            
            # 4. Удаляем временный файл после установки (опционально, можно оставить до закрытия)
            # Windows обычно блокирует файл, пока он используется как иконка, 
            # поэтому удаление может вызвать ошибку. Лучше оставить как есть, 
            # ОС сама очистит temp при перезагрузке или через время.
            
        except Exception as e:
            print(f"Не удалось установить иконку окна: {e}")
        # ======================================

        # Сетка окна
        self.grid_columnconfigure(0, weight=1)
        self.grid_columnconfigure(1, weight=2)
        self.grid_rowconfigure(0, weight=1)

        # === Левая панель (Настройки) ===
        self.left_frame = ctk.CTkFrame(self, corner_radius=10)
        self.left_frame.grid(row=0, column=0, sticky="nsew", padx=10, pady=10)
        self.left_frame.grid_columnconfigure(0, weight=1)

        # Заголовок
        self.label_title = ctk.CTkLabel(self.left_frame, text="Параметры канала", font=("Arial", 16, "bold"))
        self.label_title.grid(row=0, column=0, padx=20, pady=(20, 10), sticky="w")

        # Тип канала (радиокнопки)
        self.label_channel_type = ctk.CTkLabel(self.left_frame, text="Тип канала:")
        self.label_channel_type.grid(row=1, column=0, padx=20, pady=(15, 5), sticky="w")
        
        self.channel_type_var = ctk.StringVar(value="b|a")
        self.frame_radio_channel = ctk.CTkFrame(self.left_frame, fg_color="transparent")
        self.frame_radio_channel.grid(row=2, column=0, padx=20, pady=5, sticky="w")
        
        channel_types = ["b|a", "a|b", "a,b"]
        for i, ct_val in enumerate(channel_types):
            rb = ctk.CTkRadioButton(self.frame_radio_channel, text=f"P({ct_val})", variable=self.channel_type_var, value=ct_val, command=self.on_channel_type_change)
            rb.grid(row=0, column=i, padx=10)

        # Размер матрицы (радиокнопки)
        self.label_size = ctk.CTkLabel(self.left_frame, text="Размер алфавита (N):")
        self.label_size.grid(row=3, column=0, padx=20, pady=(15, 5), sticky="w")
        
        self.size_var = ctk.StringVar(value="3x3")
        self.frame_radio_size = ctk.CTkFrame(self.left_frame, fg_color="transparent")
        self.frame_radio_size.grid(row=4, column=0, padx=20, pady=5, sticky="w")
        
        sizes = ["3x3", "4x4", "5x5"]
        for i, size_val in enumerate(sizes):
            rb = ctk.CTkRadioButton(self.frame_radio_size, text=size_val, variable=self.size_var, value=size_val, command=self.on_size_change)
            rb.grid(row=0, column=i, padx=10)

        # Контейнер для динамических элементов
        self.dynamic_container = ctk.CTkFrame(self.left_frame, fg_color="transparent")
        self.dynamic_container.grid(row=5, column=0, padx=20, pady=10, sticky="nsew")
        self.left_frame.grid_rowconfigure(5, weight=1)

        # === Кнопки действий в 2 ряда ===
        self.btn_frame = ctk.CTkFrame(self.left_frame, fg_color="transparent")
        self.btn_frame.grid(row=6, column=0, padx=20, pady=20, sticky="ew")
        
        # Верхний ряд кнопок
        self.btn_row1 = ctk.CTkFrame(self.btn_frame, fg_color="transparent")
        self.btn_row1.pack(fill="x", pady=(0, 5))
        
        self.btn_import = ctk.CTkButton(self.btn_row1, text="📂 Импорт", command=self.import_from_excel, fg_color="#9b59b6")
        self.btn_import.pack(side="left", expand=True, fill="x", padx=5)

        self.btn_random = ctk.CTkButton(self.btn_row1, text="Случайно", command=self.fill_random, fg_color="#1f6feb")
        self.btn_random.pack(side="left", expand=True, fill="x", padx=5)
        
        # Нижний ряд кнопок
        self.btn_row2 = ctk.CTkFrame(self.btn_frame, fg_color="transparent")
        self.btn_row2.pack(fill="x")
        
        self.btn_clear = ctk.CTkButton(self.btn_row2, text="Очистить", command=self.clear_all, fg_color="#da3633")
        self.btn_clear.pack(side="left", expand=True, fill="x", padx=5)
        
        self.btn_calc = ctk.CTkButton(self.btn_row2, text="Рассчитать", command=self.calculate, fg_color="#2ea043")
        self.btn_calc.pack(side="left", expand=True, fill="x", padx=5)

        # === Правая панель (Результаты) ===
        self.right_frame = ctk.CTkFrame(self, corner_radius=10)
        self.right_frame.grid(row=0, column=1, sticky="nsew", padx=10, pady=10)
        self.right_frame.grid_columnconfigure(0, weight=1)
        self.right_frame.grid_rowconfigure(1, weight=1)

        self.label_results = ctk.CTkLabel(self.right_frame, text="Результаты", font=("Arial", 16, "bold"))
        self.label_results.grid(row=0, column=0, padx=20, pady=(15, 5), sticky="w")

        # ВАШИ ИЗМЕНЕНИЯ: Шрифт Arial 16 (как в вашем коде)
        self.text_result = ctk.CTkTextbox(self.right_frame, wrap="none", font=("Arial", 16), height=300)
        self.text_result.grid(row=1, column=0, padx=20, pady=(0, 20), sticky="nsew")
        self.text_result.configure(state="disabled", fg_color="transparent", border_width=0)

        # Инициализация
        self.current_size = 3
        self.matrix_entries = [] 
        self.vector_entries = [] 
        self.create_inputs(self.current_size, self.channel_type_var.get())

    def on_channel_type_change(self, choice=None):
        self.create_inputs(self.current_size, self.channel_type_var.get())

    def on_size_change(self, choice=None):
        size_str = self.size_var.get().split('x')[0]
        new_size = int(size_str)
        if new_size != self.current_size:
            self.current_size = new_size
            self.create_inputs(new_size, self.channel_type_var.get())

    def create_inputs(self, size, channel_type):
        """Создание полей ввода"""
        for widget in self.dynamic_container.winfo_children():
            widget.destroy()
        
        self.matrix_entries = []
        self.vector_entries = []

        # --- Матрица ---
        matrix_label_text = ""
        vector_label_text = ""

        if channel_type == "b|a":
            matrix_label_text = f"Матрица P(Bj|Ai) [{size}x{size}]:"
            vector_label_text = f"Вероятности P(Ai) [{size}]:"
        elif channel_type == "a|b":
            matrix_label_text = f"Матрица P(Ai|Bj) [{size}x{size}]:"
            vector_label_text = f"Вероятности P(Bj) [{size}]:"
        elif channel_type == "a,b":
            matrix_label_text = f"Матрица P(Ai,Bj) [{size}x{size}]:"
            vector_label_text = "" 

        lbl_matrix = ctk.CTkLabel(self.dynamic_container, text=matrix_label_text)
        lbl_matrix.grid(row=0, column=0, sticky="w", pady=(0, 5))
        
        frame_matrix = ctk.CTkFrame(self.dynamic_container, fg_color="transparent")
        frame_matrix.grid(row=1, column=0, pady=5)
        
        for i in range(size):
            row_entries = []
            for j in range(size):
                entry = ctk.CTkEntry(frame_matrix, width=50, placeholder_text="0.00")
                val = random.uniform(0.1, 1.0)
                entry.insert(0, f"{val:.2f}")
                entry.grid(row=i, column=j, padx=2, pady=2)
                row_entries.append(entry)
            self.matrix_entries.append(row_entries)

        # --- Вектор (если нужен) ---
        if channel_type != "a,b":
            lbl_vector = ctk.CTkLabel(self.dynamic_container, text=vector_label_text)
            lbl_vector.grid(row=2, column=0, sticky="w", pady=(15, 5))
            
            frame_vector = ctk.CTkFrame(self.dynamic_container, fg_color="transparent")
            frame_vector.grid(row=3, column=0, pady=5)
            
            for i in range(size):
                entry = ctk.CTkEntry(frame_vector, width=50, placeholder_text="0.00")
                val = random.uniform(0.1, 1.0)
                entry.insert(0, f"{val:.2f}")
                entry.grid(row=0, column=i, padx=2, pady=2)
                self.vector_entries.append(entry)
        else:
            self.vector_entries = []

    def import_from_excel(self):
        """Импорт данных из файла Excel по шаблону import.xlsx"""
        file_path = fd.askopenfilename(filetypes=[("Excel files", "*.xlsx *.xls")])
        if not file_path:
            return

        try:
            wb = openpyxl.load_workbook(file_path)
            sheet_names = wb.sheetnames
            
            # Логика поиска листов Matrix и Vector
            matrix_sheet_name = None
            vector_sheet_name = None
            
            if 'Matrix' in sheet_names: matrix_sheet_name = 'Matrix'
            elif 'Лист1' in sheet_names: matrix_sheet_name = 'Лист1'
            else: matrix_sheet_name = sheet_names[0]
                
            if 'Vector' in sheet_names: vector_sheet_name = 'Vector'
            elif len(sheet_names) > 1:
                 for name in sheet_names:
                     if name != matrix_sheet_name:
                         vector_sheet_name = name
                         break

            if not matrix_sheet_name: raise ValueError("Не найден лист с матрицей!")

            ws_matrix = wb[matrix_sheet_name]
            data_matrix = []
            for row in ws_matrix.iter_rows(values_only=True):
                if any(cell is not None for cell in row):
                    clean_row = [cell for cell in row if isinstance(cell, (int, float))]
                    if clean_row: data_matrix.append(clean_row)
            
            if not data_matrix: raise ValueError("Матрица пуста!")

            n = len(data_matrix)
            if n != len(data_matrix[0]): raise ValueError(f"Матрица должна быть квадратной! Получено {n}x{len(data_matrix[0])}")
            if n not in [3, 4, 5]: raise ValueError(f"Поддерживаются размеры 3x3, 4x4, 5x5. Получено {n}x{n}")

            # Обновляем интерфейс
            self.size_var.set(f"{n}x{n}")
            self.current_size = n
            self.create_inputs(n, self.channel_type_var.get())

            # Заполняем матрицу
            for i in range(n):
                for j in range(n):
                    if i < len(data_matrix) and j < len(data_matrix[i]):
                        self.matrix_entries[i][j].delete(0, "end")
                        self.matrix_entries[i][j].insert(0, str(data_matrix[i][j]))

            # Заполняем вектор
            if self.channel_type_var.get() != "a,b" and vector_sheet_name:
                ws_vector = wb[vector_sheet_name]
                vector_data = []
                for row in ws_vector.iter_rows(values_only=True):
                    if any(cell is not None for cell in row):
                        clean_row = [cell for cell in row if isinstance(cell, (int, float))]
                        vector_data.extend(clean_row)
                
                if len(vector_data) >= n:
                    for i in range(n):
                        self.vector_entries[i].delete(0, "end")
                        self.vector_entries[i].insert(0, str(vector_data[i]))
                else:
                    ctk.CTkMessagebox(title="Предупреждение", message=f"Вектор слишком короткий.", icon="warning")

            ctk.CTkMessagebox(title="Успех", message=f"Данные импортированы! Размер: {n}x{n}", icon="check")

        except Exception as e:
            ctk.CTkMessagebox(title="Ошибка импорта", message=f"Не удалось загрузить файл:\n{str(e)}", icon="cancel")

    def normalize_vector(self, vec):
        s = sum(vec)
        if s == 0: return [1/len(vec)] * len(vec)
        return [v/s for v in vec]

    def normalize_matrix_rows(self, matrix):
        return [self.normalize_vector(row) for row in matrix]

    def normalize_matrix_cols(self, matrix):
        n = len(matrix)
        if n == 0: return []
        m = len(matrix[0])
        normalized = [[0.0]*m for _ in range(n)]
        for j in range(m):
            col_sum = sum(matrix[i][j] for i in range(n))
            if col_sum > 0:
                for i in range(n):
                    normalized[i][j] = matrix[i][j] / col_sum
            else:
                for i in range(n):
                    normalized[i][j] = 1.0/n if n > 0 else 0
        return normalized

    def fill_random(self):
        channel_type = self.channel_type_var.get()
        for row in self.matrix_entries:
            for entry in row:
                entry.delete(0, "end")
                entry.insert(0, f"{random.uniform(0.1, 1.0):.2f}")
        
        if channel_type != "a,b":
            for entry in self.vector_entries:
                entry.delete(0, "end")
                entry.insert(0, f"{random.uniform(0.1, 1.0):.2f}")

    def clear_all(self):
        for row in self.matrix_entries:
            for entry in row:
                entry.delete(0, "end")
        if self.vector_entries:
            for entry in self.vector_entries:
                entry.delete(0, "end")
        self.update_result_text("")

    def get_values(self):
        try:
            raw_matrix = [[float(e.get()) for e in row] for row in self.matrix_entries]
            channel_type = self.channel_type_var.get()
            
            if channel_type == "b|a":
                raw_vector = [float(e.get()) for e in self.vector_entries]
                pa = self.normalize_vector(raw_vector)
                matrix = self.normalize_matrix_rows(raw_matrix)
                return matrix, pa, None
            elif channel_type == "a|b":
                raw_vector = [float(e.get()) for e in self.vector_entries]
                pb = self.normalize_vector(raw_vector)
                matrix = self.normalize_matrix_cols(raw_matrix)
                return matrix, None, pb
            elif channel_type == "a,b":
                total_sum = sum(sum(row) for row in raw_matrix)
                if total_sum == 0: raise ValueError("Сумма всех элементов матрицы P(A,B) не может быть 0")
                matrix = [[p / total_sum for p in row] for row in raw_matrix]
                return matrix, None, None

        except ValueError:
            ctk.CTkMessagebox(title="Ошибка", message="Введите корректные числа!", icon="cancel")
            return None, None, None

    def calculate_entropy(self, prob_vector):
        h = 0.0
        for p in prob_vector:
            if p > 1e-9:
                h -= p * math.log2(p)
        return h

    def calculate_from_joint(self, joint_matrix):
        n = len(joint_matrix)
        if n == 0: return {}

        pa = [sum(row) for row in joint_matrix]
        pb = [sum(joint_matrix[i][j] for i in range(n)) for j in range(n)]

        h_a = self.calculate_entropy(pa)
        h_b = self.calculate_entropy(pb)
        
        h_ab_direct = 0.0
        for i in range(n):
            for j in range(n):
                p_ab = joint_matrix[i][j]
                if p_ab > 1e-9:
                    h_ab_direct -= p_ab * math.log2(p_ab)
        
        h_ai_list = []
        h_b_given_a_weighted_sum = 0.0
        for i in range(n):
            if pa[i] > 1e-9:
                row_dist = [joint_matrix[i][j] / pa[i] for j in range(n)]
                h_cond = self.calculate_entropy(row_dist)
                h_ai_list.append(h_cond)
                h_b_given_a_weighted_sum += pa[i] * h_cond
            else:
                h_ai_list.append(0.0)
        
        h_b_given_a = h_b_given_a_weighted_sum

        h_bj_list = []
        h_a_given_b_weighted_sum = 0.0
        for j in range(n):
            if pb[j] > 1e-9:
                col_dist = [joint_matrix[i][j] / pb[j] for i in range(n)]
                h_cond = self.calculate_entropy(col_dist)
                h_bj_list.append(h_cond)
                h_a_given_b_weighted_sum += pb[j] * h_cond
            else:
                h_bj_list.append(0.0)
            
        h_a_given_b = h_a_given_b_weighted_sum

        return {
            'pa': pa, 'pb': pb, 
            'h_a': h_a, 'h_b': h_b, 'h_ab': h_ab_direct, 
            'h_b_given_a': h_b_given_a, 'h_a_given_b': h_a_given_b,
            'h_bj_list': h_bj_list, 'h_ai_list': h_ai_list
        }

    def calculate_uniform_conditional_entropy(self, matrix, channel_type):
        n = len(matrix)
        if n == 0: return 0
        uniform_p = 1.0 / n
        
        h_conditional_uniform = 0.0
        
        if channel_type == "b|a":
            for i in range(n):
                h_bi_given_ai = self.calculate_entropy(matrix[i])
                h_conditional_uniform += uniform_p * h_bi_given_ai
                
        elif channel_type == "a|b":
            for j in range(n):
                col_dist = [matrix[i][j] for i in range(n)]
                h_aj_given_bj = self.calculate_entropy(col_dist)
                h_conditional_uniform += uniform_p * h_aj_given_bj
                
        elif channel_type == "a,b":
            pa_orig = [sum(row) for row in matrix]
            for i in range(n):
                if pa_orig[i] > 1e-9:
                    row_dist = [matrix[i][j] / pa_orig[i] for j in range(n)]
                    h_bi_given_ai = self.calculate_entropy(row_dist)
                    h_conditional_uniform += uniform_p * h_bi_given_ai
                else:
                    h_conditional_uniform += 0.0

        return h_conditional_uniform


    def calculate(self):
        matrix, pa, pb = self.get_values()
        if not matrix: return

        n = len(matrix)
        channel_type = self.channel_type_var.get()
        results = {}

        if channel_type == "b|a":
            joint_matrix = [[pa[i] * matrix[i][j] for j in range(n)] for i in range(n)]
            results = self.calculate_from_joint(joint_matrix)
            
        elif channel_type == "a|b":
            joint_matrix = [[matrix[i][j] * pb[j] for j in range(n)] for i in range(n)]
            results = self.calculate_from_joint(joint_matrix)
            
        elif channel_type == "a,b":
            results = self.calculate_from_joint(matrix)
            joint_matrix = matrix

        h_conditional_uniform = self.calculate_uniform_conditional_entropy(matrix, channel_type)
        h_max = math.log2(n) if n > 1 else 0

        result_lines = []
        
        # ВАШИ ИЗМЕНЕНИЯ: Сокращенный вывод результатов
        if channel_type == "b|a":
            result_lines = [
                f"H(A)max: {h_max:.3f}",
                f"H(A): {results['h_a']:.3f}",
                f"H(B): {results['h_b']:.3f}",    
            ]
            # Оставляем только H(ai) и итоговые условные энтропии
            for i, h_val in enumerate(results['h_ai_list']):
                result_lines.append(f"H(a{i+1}): {h_val:.3f}")
            result_lines.extend([
                f"H(B/A) равн.: {h_conditional_uniform:.3f}",
                f"H(B/A) нерав.: {results['h_b_given_a']:.3f}"
            ])
            
        elif channel_type == "a|b":
            result_lines = [
                f"H(B)max: {h_max:.3f}",
                f"H(B): {results['h_b']:.3f}",
                f"H(A): {results['h_a']:.3f}",
            ]
            for i, h_val in enumerate(results['h_bj_list']):
                result_lines.append(f"H(b{i+1}): {h_val:.3f}")
            result_lines.extend([
                f"H(A/B) равн.: {h_conditional_uniform:.3f}",
                f"H(A/B) нерав.: {results['h_a_given_b']:.3f}"
            ])
            
        elif channel_type == "a,b":
            result_lines = [
                f"H(B)max: {h_max:.3f}",
                f"H(B): {results['h_b']:.3f}",
                f"H(A): {results['h_a']:.3f}",
                f"H(A,B): {results['h_ab']:.3f}",
            ]
            # Для a,b выводим H(ai) и H(bi) как в Excel
            for i, h_val in enumerate(results['h_ai_list']):
                result_lines.append(f"H(a{i+1}): {h_val:.3f}")
            
            for i, h_val in enumerate(results['h_bj_list']):
                result_lines.append(f"H(b{i+1}): {h_val:.3f}")
            
            result_lines.extend([
                f"H(B/A) нерав.: {results['h_b_given_a']:.3f}",
                f"H(A/B) нерав.: {results['h_a_given_b']:.3f}"
            ])

        self.update_result_text("\n".join(result_lines))

    def update_result_text(self, text):
        self.text_result.configure(state="normal")
        self.text_result.delete("1.0", "end")
        if text:
            self.text_result.insert("1.0", text)
        self.text_result.configure(state="disabled")

if __name__ == "__main__":
    app = EntropyCalculatorApp()
    app.mainloop()